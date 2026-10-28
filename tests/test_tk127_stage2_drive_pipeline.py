"""Synthetic indexed retrieval, policy, and all four real report exporters."""
import copy
import io
import json
from types import SimpleNamespace

import openpyxl
import pytest
from docx import Document
from pypdf import PdfReader

from packages.common.enums import ExternalAIPolicy, VerificationProfile
from packages.common.schemas import Block, NormalizedDocument, Page
from packages.rag_engine.review import review_document
from packages.verification_engine.pipeline import DocumentResult, ProjectContext
from packages.verification_engine.review_items import build_document_review_items
from test_drive_rag import FILE, library, metadata  # noqa: F401
from test_report_completion import report_case, create  # noqa: F401
from test_tk127_stage2_reference_binding import make_document, pilot_cases


def test_legacy_frozen_finding_views_are_not_re_evaluated():
    from packages.report_engine.reference_opinions import reference_rows

    for sample in [0, 10, 20]:
        document = make_document(pilot_cases()[sample])
        document.findings = [SimpleNamespace(title="합성 고정본 표시 전용 지적")]
        assert reference_rows(document)[0]["reference_observations"]
        document.engine_data = {}
        assert reference_rows(document) == []


@pytest.mark.parametrize("index", [1, 2, 3])
@pytest.mark.parametrize("policy", [ExternalAIPolicy.MASKED, ExternalAIPolicy.ORIGINAL])
def test_direct_reference_preserves_external_transmission_policy(library, monkeypatch, index, policy):
    from packages.llm_router.router import RouterResult
    from packages.pii_engine.engine import PIIEngine

    lib, drive = library
    email = f"synthetic{index}@example.invalid"
    number = f"2099다55{index:03d}"
    text = f"{number} 합성 참고문헌 연락처 {email}. 보관 의무를 검토합니다."
    drive.contents[FILE] = text
    drive.items[FILE] = {**metadata(text=text), "name": email + ".txt"}
    lib.sync()
    monkeypatch.setattr(lib, "select", lambda _: {"sources": [], "decision": "NOT_USED", "reason": "NO_SHARED_KEY_TERMS"})
    result = DocumentResult("d", "합성.txt", normalized=NormalizedDocument("d", "합성.txt", "text/plain", "synthetic",
        pages=[Page(1, blocks=[Block("b", text, 1)])]))
    result.citations = [{"citation_id": "c", "case_number": number, "raw_text": number}]
    sent = []
    class Router:
        def has_available_provider(self, **kwargs):
            return True
        async def run(self, role, request, **kwargs):
            sent.append((request.user, kwargs["policy"]))
            return RouterResult(used=True, parsed={"observations": []})
    pii = PIIEngine(SimpleNamespace(pseudonym_for=lambda kind, text: "합성가림"))
    context = ProjectContext("p", external_ai_policy=policy)
    review_document(result, lib, Router(), context, pii)
    assert sent and all(p == policy for _, p in sent)
    assert all((email in text) == (policy == ExternalAIPolicy.ORIGINAL) for text, _ in sent)


@pytest.mark.parametrize("index", [1, 2, 3])
@pytest.mark.parametrize("profile,policy", [(VerificationProfile.QUICK, ExternalAIPolicy.ORIGINAL),
                                            (VerificationProfile.STANDARD, ExternalAIPolicy.LOCAL_ONLY)])
def test_indexed_citation_lookup_and_no_external_model(library, monkeypatch, index, profile, policy):
    lib, drive = library
    number = f"2098나77{index:03d}"
    text = f"{number} 합성 판결의 자료 보관에 관한 참고 설명입니다. " * 2
    drive.contents[FILE] = text
    drive.items[FILE] = metadata(text=text)
    lib.sync()
    citation = {"citation_id": "c", "case_number": number, "raw_text": number}
    assert lib.search_citation(citation)[0]["text"] in text
    assert not lib.search_citation({**citation, "case_number": f"2098나66{index:03d}"})
    # Simulate document-level top-12 omitting this source. Direct indexed query
    # still executes; no replacement of the search_citation implementation.
    monkeypatch.setattr(lib, "select", lambda _: {"sources": [], "decision": "NOT_USED", "reason": "NO_SHARED_KEY_TERMS"})
    normalized = NormalizedDocument("d", "합성서면.txt", "text/plain", "synthetic",
                                    pages=[Page(1, blocks=[Block("b", number + " 자료 보관에 관한 주장입니다.", 1)])])
    result = DocumentResult("d", "합성서면.txt", normalized=normalized)
    result.citations = [citation]
    class NoModel:
        def has_available_provider(self, **kwargs):
            raise AssertionError("local branch must short-circuit provider lookup")
        async def run(self, *args, **kwargs):
            raise AssertionError("external model called")
    context = ProjectContext("p", profile=profile, external_ai_policy=policy)
    review = review_document(result, lib, NoModel(), context, SimpleNamespace(mask_text=lambda s: SimpleNamespace(masked_text=s)))
    assert review["observations"] and not review["model_executed"]
    assert build_document_review_items(result)[0].to_dict()["reference_status"] == "SUPPORTED"
    drive.items.clear()
    lib.sync()
    assert not lib.search_citation(citation)


@pytest.mark.parametrize("index", [0, 10, 20])
def test_four_reports_retain_unofficial_reference_opinions(report_case, index):
    case = report_case
    synthetic = make_document(pilot_cases()[index])
    doc = vars(synthetic)
    doc["document_id"] = "d1"
    doc["filename"] = "합성서면.pdf"
    doc["review_items"] = [i.to_dict() for i in build_document_review_items(SimpleNamespace(**doc))]
    source_quote = doc["engine_data"]["rag"]["observations"][0]["source_quote"]
    result = copy.deepcopy(case.run.result_json)
    result["documents"] = [doc]
    case.run.result_json = result
    case.session.delete(case.finding)
    case.session.commit()
    report = create(case, formats=["pdf", "docx", "xlsx", "json"])
    blobs = {kind: case.client.get(f"/api/reports/{report['report_id']}/download/{kind}").content for kind in report["artifacts"]}
    # Check real exported cell contents, not helper strings alone.
    pdf_text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(blobs["pdf"])).pages)
    docx = Document(io.BytesIO(blobs["docx"]))
    cells = [c.text for table in docx.tables for row in table.rows for c in row.cells]
    assert any(source_quote in c and "공식 확인 아님" in c for c in cells)
    assert "공식 확인 아님" in pdf_text and source_quote.split()[0] in pdf_text
    workbook = openpyxl.load_workbook(io.BytesIO(blobs["xlsx"]))
    sheet = workbook["참고문헌대조"]
    assert "공식 확인" in sheet.cell(1, 1).value
    assert sheet.cell(2, 5).value == "공식 여부"
    assert "공식 확인 아님" == sheet.cell(3, 5).value
    assert source_quote in sheet.cell(3, 3).value
    payload = json.loads(blobs["json"])
    observation = payload["documents"][0]["reference_observations"][0]
    assert observation["source"] == "google_drive"
    assert observation["official"] is False and observation["binding"] is False
