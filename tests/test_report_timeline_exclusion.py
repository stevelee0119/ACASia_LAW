"""TK-86: 보고서 외부 산출물 timeline 제외 및 실행 결과 timeline 보존 서버 검증 시험.

수용 요건:
1. 실행 결과 to_payload()에는 timeline이 유지된다 (엔진 분석 결과 및 채점기 입력 무결성).
2. 보고서 JSON(to_payload with report_metadata), xlsx 부록(to_xlsx), 기술 스냅숏(technical_payload)에는 timeline이 없다.
"""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

import openpyxl

from packages.common.enums import JobState
from packages.report_engine.exporters import to_payload, to_xlsx
from packages.report_engine.snapshot import technical_payload


def test_timeline_retained_in_execution_payload_but_excluded_from_reports():
    """실행 결과 to_payload에는 timeline이 보존되고, 보고서/스냅숏/xlsx에는 timeline이 제외된다."""
    # 사건 일지(timeline) 이벤트 데이터를 포함하는 모의 실행 결과 객체 생성
    mock_timeline = [
        {"date": "2023-01-15", "event": "계약 체결", "source_doc": "계약서.pdf", "page": 1},
        {"date": "2023-03-20", "event": "중도금 지급", "source_doc": "영수증.pdf", "page": 1},
    ]
    run_result = SimpleNamespace(
        run_id="run_test_123",
        project_id="proj_test_456",
        verification_key="vkey_abc",
        state=JobState.COMPLETED,
        started_at=datetime(2026, 10, 20, 10, 0, 0),
        finished_at=datetime(2026, 10, 20, 10, 1, 0),
        scores={"overall": 95},
        timeline=mock_timeline,
        unavailable_sources=[],
        unverified_items=[],
        errors=[],
        findings=[],
        project_findings=[],
        all_findings=[],
        model_executions=[],
        run_manifest={},
        documents=[],
    )

    # 1. 실행 결과 to_payload: timeline 필드가 정상 보존되어야 한다
    exec_payload = to_payload(run_result)
    assert "timeline" in exec_payload, "실행 결과 to_payload에는 timeline이 존재해야 한다"
    assert exec_payload["timeline"] == mock_timeline

    # 2. 기술 스냅숏 technical_payload: timeline 필드가 제외되어야 한다
    snapshot = technical_payload(run_result)
    assert "timeline" not in snapshot, "technical_payload 루트에 timeline이 없어야 한다"
    assert "timeline" not in snapshot.get("review_snapshot", {}), "technical_payload review_snapshot에 timeline이 없어야 한다"

    # 3. 보고서 결과 to_payload (report_metadata 존재 시): timeline이 제거되어야 한다
    run_result.report_metadata = {
        "id": "rep_test_789",
        "title": "검증 보고서",
        "created_at": "2026-10-20T10:05:00",
    }
    run_result.review_snapshot = {"engine_result": {"scores": {"overall": 95}}}

    report_payload = to_payload(run_result)
    assert "timeline" not in report_payload, "보고서 payload에는 timeline이 없어야 한다"

    # 4. Excel 기술부록 to_xlsx: 기술부록 시트에 timeline 경로가 없어야 한다
    xlsx_bytes = to_xlsx(run_result)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=True)
    assert "기술부록" in wb.sheetnames
    tech_sheet = wb["기술부록"]
    cell_values = [str(cell.value) for row in tech_sheet.iter_rows() for cell in row if cell.value is not None]
    assert not any("timeline" in val for val in cell_values), "Excel 기술부록 시트에 timeline 경로가 나타나지 않아야 한다"


def test_reports_api_snapshot_excludes_timeline():
    """reports API 내부 _snapshot 생성 시 engine_result에서 timeline이 제외되고 원본은 보존된다."""
    from apps.api.routers.reports import _snapshot

    mock_project = SimpleNamespace(scope_revision=1)
    mock_run = SimpleNamespace(id="run_1", scope_revision=1)
    mock_engine = {
        "timeline": [{"date": "2023-01-01", "event": "이벤트", "source_doc": "계약서.pdf"}],
        "scores": {"overall": 90},
        "errors": [],
        "findings": [],
        "documents": [],
    }
    mock_metadata = {
        "audience": "INTERNAL",
        "state": "DRAFT",
        "created_by": "tester",
        "created_at": "2026-10-21T10:00:00",
    }
    mock_human = {
        "workflow": [],
        "matrix": {"claims": []},
        "case_profile": {},
        "review_history": [],
    }
    mock_manifest = {}

    snapshot = _snapshot(None, mock_project, mock_run, mock_engine, mock_metadata, mock_manifest, human=mock_human)

    # 1. snapshot 내부 engine_result에서 timeline이 제거되었는지 확인
    assert "timeline" not in snapshot["engine_result"], "스냅숏의 engine_result에는 timeline이 없어야 한다"
    # 2. 원본 engine 객체의 timeline은 유지되는지 확인 (부수효과 방지)
    assert "timeline" in mock_engine, "원본 engine 딕셔너리의 timeline은 보존되어야 한다"

