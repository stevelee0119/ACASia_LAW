"""Automatic source proofs, baseline preservation, and actual-format controls."""
import json
import pytest
from packages.audit_engine import AuditChain
from packages.common.storage import sha256_file
from packages.document_engine.reading_text import build_reading_text
from packages.legal_engine.defense_source_delta import source_defense_proof, SOURCE_SENSES
from packages.legal_engine.legal_rules import load_defense_groups
from packages.verification_engine import DocumentInput, ProjectContext, VerificationPipeline
from test_tk45_source_negative_delta_spike import rows, RULE


def run_source(path, body, format='txt', width=595):
    source='청구원인\n'+body
    if format=='txt':
        path.write_text(source,encoding='utf-8');mime='text/plain'
    elif format=='docx':
        from docx import Document
        doc=Document();doc.add_paragraph('청구원인');doc.add_paragraph(body);doc.save(path)
        mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    else:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfgen.canvas import Canvas
        name='HYSMyeongJo-Medium';pdfmetrics.registerFont(UnicodeCIDFont(name))
        canvas=Canvas(str(path),pagesize=(width,842));canvas.setFont(name,12);canvas.drawString(36,794,'청구원인')
        canvas.setFont(name,4);canvas.drawString(36,758,body);canvas.save();mime='application/pdf'
    result=VerificationPipeline(audit=AuditChain()).run('source-'+path.stem,ProjectContext('public'),[
        DocumentInput('generated',str(path),path.name,mime,sha256_file(path))])
    assert not result.errors
    text=build_reading_text(result.documents[0].normalized).text
    assert body in text
    findings=[f for f in result.all_findings if RULE in f.tags]
    return result,text,findings


def check_proof(finding,text,kind):
    proof=finding.confidence_features['source_defense_proof']
    assert proof['kind']==kind
    assert proof['selected_sense']['id'] in SOURCE_SENSES
    assert proof['selected_sense']['definition']==SOURCE_SENSES[proof['selected_sense']['id']]['definition']
    assert proof['authored_speaker']=='IMPLICIT_WRITER'
    assert proof['speaker_scope']=='FIRST_NONEMPTY_BODY_SENTENCE_ONLY'
    assert proof['authored_actor']=='UNSPECIFIED_WITHIN_CLOSED_LOCAL_CHAIN'
    assert proof['normative_scope']=='SAME_SENTENCE_EXPLICIT_CIVIL_DEFENSE'
    assert proof['application_role'] in ('ASSERTED_DIRECT_ACT_DENIAL','ASSERTED_CATEGORICAL_LEGAL_EFFECT')
    for key in ['supports_asserted_fulfillment','total_requirement_completion','facts_verified','legal_sufficiency_verified','baseline_observations_changed','actor_entity_identity_verified','acts_merged','premise_assertion_verified']:
        assert proof[key] is False
    offset=proof['reading_text_offset']
    for key,source in proof['source_slices'].items():
        a,b=proof['source_spans'][key];assert text[offset+a:offset+b]==source
    assert proof['source_slices']['predicate']
    assert proof['source_slices']['civil_norm']=='민법 제487조'
    assert 'authored_scope' not in proof['source_slices']
    assert finding.span==(offset,offset+len(finding.confidence_features['claim']))
    assert str(finding.status)=='SUSPICIOUS'
    assert str(finding.evidence_grade)=='B'
    assert '과대주장' in finding.confidence_features['verdict'] or '범주적' in finding.confidence_features['verdict']
    assert proof['accounting']['whole_document_prefix_searches']==0
    assert proof['accounting']['requirement_prefix_queries']<=2*len(load_defense_groups()['structural_signals']['stated_requirements'])
    assert proof['accounting']['conclusion_prefix_queries']<=len(load_defense_groups()['structural_signals']['limited_conclusions'])
    assert json.loads(json.dumps(finding.to_dict(),ensure_ascii=False))['confidence_features']['source_defense_proof']==proof


@pytest.mark.parametrize('name,body,relation,desired',rows,ids=[r[0] for r in rows])
def test_enabled_actual_txt_adds_only_declared_source_proofs(tmp_path,monkeypatch,name,body,relation,desired):
    with monkeypatch.context() as patch:
        patch.setattr('packages.legal_engine.legal_rules.source_defense_proof',lambda *args:None)
        before_result,_,baseline=run_source(tmp_path/'enabled.txt',body)
    result,text,findings=run_source(tmp_path/'enabled.txt',body)
    # Citation IDs are freshly generated per run, not semantic verdicts.
    unchanged=lambda r:[(f.engine,f.type,f.status,f.severity,f.evidence_grade,f.title,
                         {k:v for k,v in f.confidence_features.items() if k!='citation_id'})
                        for f in r.all_findings if RULE not in f.tags]
    assert unchanged(before_result)==unchanged(result)
    assert len(findings)==(1 if desired or relation=='PRESERVE_EXISTING_WARNING' else 0)
    if desired:
        assert baseline==[]
        kind='DIRECT_REQUIREMENT_DENIAL' if relation=='DIRECT_CASE_BOUND_DENIAL' else 'LEGAL_EFFECT_EXTENSION'
        check_proof(findings[0],text,kind)
        # The later pipeline consumes the existing legal-rule finding and
        # does not add a second judgment for that sentence.
        assert sum(RULE in f.tags for f in result.all_findings)==1
    else:
        for before,after in zip(baseline,findings):
            assert before.status==after.status and before.evidence_grade==after.evidence_grade
            assert before.confidence_features==after.confidence_features
            assert 'source_defense_proof' not in after.confidence_features


DIRECT='설령 계약이 유효하더라도, 민법 제487조의 변제공탁을 하지 않았으나 이 사건 채무에 관한 책임을 질 수 없다.'
EFFECT='설령 계약이 인정되더라도, 민법 제487조에 따라 변제공탁을 완료하였으므로 행정제재를 받지 않는다.'
OPPOSITE=[
 DIRECT.replace('하지 않았으나','하지 않았다면'),
 DIRECT.replace('하지 않았으나','하지 않은 것은 아니므로'),
 DIRECT.replace('하지 않았으나','하지 않을 계획이므로'),
 DIRECT.replace('하지 않았으나','할 필요가 없으므로'),
 DIRECT.replace('변제공탁을','일부 변제공탁을'),
 DIRECT.replace('변제공탁을 하지','변제공탁을 부분적으로 하지'),
 DIRECT.replace('변제공탁을 하지','변제공탁 서류를 작성하지'),
 DIRECT.replace('민법 제487조의','민법 제487조의 피고는'),
 DIRECT.replace('변제공탁을','원고의 변제공탁을'),
 DIRECT.replace('민법 제487조의','민법 제487조의 별건'),
 DIRECT.replace('하지 않았으나','하지 않았으나 변제공탁을 완료하였으므로'),
 DIRECT.replace('하지 않았으나','하지 않았으나 상계의 의사표시를 하였으므로'),
 DIRECT.replace('민법 제487조의','민법 제487조의 절차가 유효하다면,'),
 DIRECT.replace('하지 않았으나','하지 않았으나 원고는'),
 DIRECT.replace('없다.','없다고 한다.'),
 '“'+DIRECT+'”라고 피고는 주장하지만 원고는 부인한다.',
 EFFECT.replace('행정제재를','소포를'),
 EFFECT.replace('행정제재를','우편물을'),
 EFFECT.replace('행정제재를','법률상 의미가 불명확한 제재를'),
 EFFECT.replace('행정제재를','행정제재(題材)를'),
 EFFECT.replace('행정제재를','일부 행정제재를'),
 EFFECT.replace('받지 않는다','받지 않았다'),
 EFFECT.replace('받지 않는다','받지 않을 것이다'),
 EFFECT.replace('받지 않는다','받지 않은 것은 아니다'),
 EFFECT.replace('받지 않는다','받을 필요가 없다'),
 EFFECT.replace('받지 않는다','받지 않을 권리가 있다'),
 EFFECT.replace('받지 않는다','면제신청한다'),
 EFFECT.replace('받지 않는다','면제 신청한다'),
 EFFECT.replace('받지 않는다','면담한다'),
 EFFECT.replace('행정제재를','원고는 행정제재를'),
 EFFECT.replace('완료하였으므로','완료하였다면'),
 EFFECT.replace('완료하였으므로','완료하였다고 보고되므로'),
 EFFECT.replace('완료하였으므로','일부만 완료하였으므로'),
 EFFECT.replace('완료하였으므로','계획을 완료하였으므로'),
 EFFECT.replace('민법 제487조에 따라','알 수 없는 자료에 따라'),
 EFFECT.replace('민법 제487조에 따라','민법 제487조에 따라 원고는'),
 EFFECT.replace('행정제재를','행정제재 서류를'),
]

@pytest.mark.parametrize('body',OPPOSITE,ids=[str(n) for n in range(len(OPPOSITE))])
def test_opposite_source_scopes_never_certify_and_keep_baseline(tmp_path,monkeypatch,body):
    assert source_defense_proof(body,load_defense_groups()) is None
    with monkeypatch.context() as patch:
        patch.setattr('packages.legal_engine.legal_rules.source_defense_proof',lambda *args:None)
        _,_,baseline=run_source(tmp_path/'baseline.txt',body)
    _,_,findings=run_source(tmp_path/'enabled.txt',body)
    assert [(f.status,f.confidence_features) for f in findings]==[(f.status,f.confidence_features) for f in baseline]


@pytest.mark.parametrize('format,width',[('txt',595),('docx',595),('pdf',595),('pdf',720)])
@pytest.mark.parametrize('body,kind',[(DIRECT,'DIRECT_REQUIREMENT_DENIAL'),(EFFECT,'LEGAL_EFFECT_EXTENSION'),(OPPOSITE[6],None)])
def test_actual_format_width_source_attachment_bidirectional(tmp_path,format,width,body,kind):
    _,text,findings=run_source(tmp_path/('public.'+format),body,format,width)
    if kind:
        assert len(findings)==1;check_proof(findings[0],text,kind)
    else:
        assert not findings


@pytest.mark.parametrize('predicate,sense',[('받지 않습니다','받다1:10'),('면한다','면하다1:2'),('면합니다','면하다1:2')])
def test_strict_finite_source_sense_and_polite_forms(tmp_path,predicate,sense):
    body=EFFECT.replace('받지 않는다',predicate)
    _,text,findings=run_source(tmp_path/'finite.txt',body)
    assert len(findings)==1;check_proof(findings[0],text,'LEGAL_EFFECT_EXTENSION')
    assert findings[0].confidence_features['source_defense_proof']['selected_sense']['id']==sense


@pytest.mark.parametrize('focus',['하지는','하지도'])
def test_focus_binds_to_direct_action_not_a_new_subject(tmp_path,focus):
    body=DIRECT.replace('하지',focus)
    _,text,findings=run_source(tmp_path/'focus.txt',body)
    assert len(findings)==1;check_proof(findings[0],text,'DIRECT_REQUIREMENT_DENIAL')


# The first implementation cannot certify inherited speaker scope. Any prior
# BODY statement keeps the new proof UNKNOWN; the original rule still runs.
INHERITED_SCOPE=[
 '피고 주장의 인용문은 다음과 같다. 「이는 피고의 설명이다. '+DIRECT+' 이는 인용문의 끝이다.」 원고는 이를 부인한다.',
 '다음 두 문장은 피고의 주장이다. '+DIRECT+' 원고는 위 주장을 반박한다.',
 '조사자가 아래 주장을 그대로 보고하였다. '+EFFECT+' 원고는 이러한 결론에 동의하지 않는다.',
 '다음 주장은 옳지 않다. '+DIRECT+' 위 주장은 법정 요건과 모순된다.',
 '이 항목은 진행 경과를 설명한다. '+DIRECT,
 '새 항목에서 작성자의 결론을 제시한다. '+EFFECT,
]

@pytest.mark.parametrize('body',INHERITED_SCOPE+[DIRECT.replace('민법 제487조','헌법 제487조'),DIRECT.replace('민법 제487조','헌법 제23조'),DIRECT.replace('민법 제487조','형법 제487조')])
def test_inherited_unknown_speaker_and_non_civil_norm_preserve_baseline(tmp_path,monkeypatch,body):
    with monkeypatch.context() as patch:
        patch.setattr('packages.legal_engine.legal_rules.source_defense_proof',lambda *args:None)
        _,_,baseline=run_source(tmp_path/'prior-baseline.txt',body)
    _,_,findings=run_source(tmp_path/'prior-enabled.txt',body)
    assert [(f.status,f.confidence_features) for f in findings]==[(f.status,f.confidence_features) for f in baseline]
    assert not any('source_defense_proof' in f.confidence_features for f in findings)


@pytest.mark.parametrize('tail',['다.','습니다.','어요!','다  .  ','다'])
def test_terminal_grouping_preserves_observed_endings(tmp_path,tail):
    body=DIRECT.removesuffix('다.')+tail
    assert source_defense_proof(body,load_defense_groups()) is not None
    # Reading text trims outer paragraph whitespace, preserving internal gaps.
    body=body.rstrip()
    _,text,findings=run_source(tmp_path/'terminal.txt',body)
    assert len(findings)==1
    check_proof(findings[0],text,'DIRECT_REQUIREMENT_DENIAL')


@pytest.mark.parametrize('shape',['trailing-whitespace','single-token','condition-token'])
def test_long_unknown_tail_never_creates_a_source_proof(shape):
    if shape=='trailing-whitespace':
        body=DIRECT.removesuffix('.')+' '*32000+'x'
    else:
        body=DIRECT.replace('하지 않았으나','가'*32000+('면' if shape=='condition-token' else '')+' 하지 않았으나')
    assert source_defense_proof(body,load_defense_groups()) is None


SOURCE_PREFIX_SCOPE=[
 '「청구원인\n'+DIRECT+' 이는 인용문의 끝이다.」',
 '다음 두 문장은 피고의 주장이다.\n청구원인\n'+DIRECT,
 '조사자가 아래 주장을 보고하였다.\n청구원인\n'+EFFECT,
 '다음 주장은 옳지 않다.\n청구원인\n'+DIRECT,
 '다른 항목의 경과를 설명한다.\n청구원인\n'+EFFECT,
 '청구취지\n청구를 기각한다.\n청구원인\n'+DIRECT,
]

@pytest.mark.parametrize('body',SOURCE_PREFIX_SCOPE)
def test_earlier_source_statement_or_open_scope_before_heading_keeps_baseline(tmp_path,monkeypatch,body):
    path=tmp_path/'source-prefix.txt';path.write_text(body,encoding='utf-8')
    def findings():
        result=VerificationPipeline(audit=AuditChain()).run('source-prefix',ProjectContext('public'),[
            DocumentInput('generated',str(path),path.name,'text/plain',sha256_file(path))])
        assert not result.errors
        return [f for f in result.all_findings if RULE in f.tags]
    with monkeypatch.context() as patch:
        patch.setattr('packages.legal_engine.legal_rules.source_defense_proof',lambda *args:None)
        baseline=findings()
    actual=findings()
    assert [(f.status,f.confidence_features) for f in actual]==[(f.status,f.confidence_features) for f in baseline]
    assert not any('source_defense_proof' in f.confidence_features for f in actual)
