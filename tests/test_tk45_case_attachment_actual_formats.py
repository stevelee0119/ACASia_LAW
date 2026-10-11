"""Small generated PDF/DOCX contrasts using the approved public TXT text."""
import pytest

from packages.audit_engine import AuditChain
from packages.common.storage import sha256_file
from packages.verification_engine import DocumentInput, ProjectContext, VerificationPipeline
from tests.test_tk45_authored_predicate_scope import LIMIT, RULE
from tests.test_tk45_clause_effect_contract import LEAD


@pytest.mark.parametrize('first,expected', [
    ('상계의 의사표시를 완료한 이후', False),
    ('상계의 의사표시를 일부 완료한 이후', True),
])
@pytest.mark.parametrize('kind,width', [('pdf', 595), ('pdf', 720), ('docx', None)])
def test_public_original_attachment_survives_actual_format_pipeline(tmp_path, first, expected, kind, width):
    text = LEAD + first + ' 변제공탁을 완료하였으므로' + LIMIT
    path = tmp_path / ('attachment.' + kind)
    if kind == 'pdf':
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfgen import canvas

        font = 'HYSMyeongJo-Medium'
        try:
            pdfmetrics.getFont(font)
        except KeyError:
            pdfmetrics.registerFont(UnicodeCIDFont(font))
        pdf = canvas.Canvas(str(path), pagesize=(width, 842))
        pdf.setFont(font, 12)
        pdf.drawString(36, 842 - 48, '청구원인')
        pdf.setFont(font, 6)
        pdf.drawString(36, 842 - 84, text)
        pdf.save()
        mime = 'application/pdf'
    else:
        from docx import Document

        docx = Document()
        docx.add_heading('청구원인', level=1)
        docx.add_paragraph(text)
        docx.save(path)
        mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    item = DocumentInput('public-attachment', str(path), path.name, mime, sha256_file(path))
    result = VerificationPipeline(audit=AuditChain()).run(
        'format-' + tmp_path.name, ProjectContext('public-attachment'), [item])
    assert not result.errors
    normalized = result.documents[0].normalized
    assert normalized
    reconstructed = '\n'.join(block.text for page in normalized.pages for block in page.blocks)
    assert text in reconstructed
    findings = [row for row in result.documents[0].findings if RULE in row.tags]
    assert bool(findings) is expected
    if expected:
        assert all(row.status.value == 'SUSPICIOUS' for row in findings)
        assert all('요건 확인 요청' in row.confidence_features['verdict'] for row in findings)
