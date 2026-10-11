"""Unapproved dated-record/EOF pending review; original product is unchanged."""
import copy
import importlib.util
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

base=Path(__file__).parent
spec=importlib.util.spec_from_file_location('tiers',base/'TK-73_post013_confirmation_tier_spike.py')
t=importlib.util.module_from_spec(spec);spec.loader.exec_module(t)
s=t.s


def source_eof(path,text,owners):
    """Original decoded-source exhaustion; never semantic heading closure."""
    original=[];extension=path.suffix
    if extension=='.txt':
        raw=path.read_text(encoding='utf-8');absolute=0
        for line_index,line in enumerate(raw.splitlines(keepends=True)):
            original.extend((c,('txt_codepoint',line_index,absolute+i)) for i,c in enumerate(line.rstrip('\r\n')) if not c.isspace())
            absolute+=len(line)
    elif extension=='.docx':
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        with ZipFile(path) as archive:
            if any(re.match(r'word/(?:header|footer|footnotes|endnotes)',name) for name in archive.namelist()):return False
            tree=ET.fromstring(archive.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        body=tree.find('w:body',ns)
        if body is None or any(node.tag.rsplit('}',1)[-1] not in ('p','sectPr') for node in body):return False
        barriers={'drawing','object','fldSimple','instrText','delText','tab','br','txbxContent','altChunk'}
        if any(node.tag.rsplit('}',1)[-1] in barriers for node in body.iter()):return False
        for p_index,p in enumerate(body.findall('w:p',ns)):
            for run_index,node in enumerate(p.findall('.//w:t',ns)):
                original.extend((c,('ooxml_text_codepoint',p_index,run_index,i)) for i,c in enumerate(node.text or '') if not c.isspace())
    else:
        import pdfplumber
        with pdfplumber.open(path) as raw:
            for page_number,page in enumerate(raw.pages,1):
                for index,char in enumerate(page.chars):
                    original.extend((c,('pdf_extraction',page_number,index,sub)) for sub,c in enumerate(char['text']) if not c.isspace())
    shown=[(c,owner[4] if owner else None) for c,owner in zip(text,owners) if not c.isspace()]
    return bool(original) and original==shown


def record_binding(payload,row):
    """Prove parser projection while retaining the full original record."""
    found=t.DATE_RE.search(payload)
    prefix=payload[:found.start()].strip() if found else payload.strip()
    date=found.group(0) if found else ''
    suffix=payload[found.end():].strip() if found else ''
    valid=(t.literal(row.get('name'))==t.literal(prefix) and
           t.literal(row.get('date'))==t.literal(date) and
           t.literal(row.get('author'))==t.literal(suffix) and
           t.literal(row.get('purpose'))==t.literal(suffix))
    return {'valid':valid,'complete_payload':payload,'canonical_name':row.get('name'),
            'parsed_date':row.get('date'),'original_author':row.get('author'),'original_purpose':row.get('purpose'),
            'date_span':list(found.span()) if found else None,
            'all_date_spans':[list(m.span()) for m in t.DATE_RE.finditer(payload)],
            'prefix':prefix,'suffix':suffix}


def diagnose(doc,path,certificates=None,missing_owner=False):
    rows=s.exhibit_rows(doc);before=copy.deepcopy(rows);findings=s.check_exhibits(doc)
    old=[f for f in findings if f.confidence_features.get('rule_id')=='EVI.EVIDENCE_NUMBER_DUPLICATE']
    reading=s.build_reading_text(doc);text=reading.text
    owners=s.source_index(reading,s.native_pieces(doc,path)) if not certificates else []
    if missing_owner:
        for i in range(len(owners)-1,-1,-1):
            if owners[i] and not text[i].isspace():owners[i]=None;break
    eof=source_eof(path,text,owners) if owners else False
    mask=s.protection(text);paragraphs=[];stack=[];active=None;offset=0
    for raw in text.split('\n'):
        start=offset+len(raw)-len(raw.lstrip());value=raw.strip();offset+=len(raw)+1
        head=s.HEAD.fullmatch(value)
        if head and not s.parse_exhibits(value) and owners and s.native(text,owners,start,start+len(value)) and not any(mask[start:start+len(value)]):
            rank=0 if head['marker'][0].isdigit() else 1
            while stack and stack[-1]['rank']>=rank:stack.pop()['end']=len(paragraphs)
            active={'index':len(paragraphs),'rank':rank,'parent':tuple(h['index'] for h in stack),'end':None,
                    'inventory':head['title'].endswith('목록') and not s._EXHIBIT_SECTION_HEAD_RE.search(value),'witnesses':set(),'open':True}
            stack.append(active)
        paragraphs.append({'start':start,'text':value,'heading':active})
    occurrences={}
    for index,p in enumerate(paragraphs):
        h=p['heading'];value=p['text'];start=p['start'];refs=s.parse_exhibits(value)
        if not h or not refs or refs[0]['span'][0]!=0:continue
        unique={tuple(r['span']):r for r in refs};spans=list(unique)
        for n,span in enumerate(spans):
            ref=unique[span];left,right=start+span[0],start+span[1]
            end=start+(spans[n+1][0] if n+1<len(spans) else len(value));tail=text[right:end]
            mark=re.match(r'\s*:\s*',tail);content=right+(mark.end() if mark else len(tail)-len(tail.lstrip()))
            good=bool(owners) and s.native(text,owners,left,right) and s.covered(text,owners,right,end)
            protected=any(mask[left:end]) or any(c in s.BRACKETS or c in s.BRACKETS.values() for c in text[left:end])
            uncertain=any(unicodedata.category(c) in ('Cf','Co','Cs') for c in text[left:end])
            good=good and not protected and not uncertain
            definition=bool(mark) and good and s.native(text,owners,right,content)
            if h['inventory']:
                if h['open'] and definition:h['witnesses'].add((left,right))
                else:h['open']=False
            colon=any(text[i] in ':：' and not mask[i] and s.native(text,owners,i-1,i+2) for i in range(content+1,end-1)) if owners else False
            record={'heading':h,'complete_payload':text[content:end].strip(),'source_complete':good,'colon_shape':colon,
                    'reference_span':(left,right),'payload_span':(content,end)}
            record['native_date_occurrences']=[s.native(text,owners,content+m.start(),content+m.end())
                                               for m in t.DATE_RE.finditer(record['complete_payload'])] if owners else []
            for branch in ref['branches'] or [None]:occurrences.setdefault((index,span,(ref['party'],ref['number'],branch)),[]).append(record)
    groups={}
    for i,row in enumerate(rows):
        for branch in row['branches'] or [None]:
            identity=(row['party'],row['number'],branch);group=row.get('line_group')
            matched=occurrences.get((group[0],tuple(group[1]),identity),[]) if group else []
            record=matched[0] if len(matched)==1 else None
            if record:record={**record,'binding':record_binding(record['complete_payload'],row)}
            groups.setdefault(identity,[]).append((i,row,record))
    tiers=[]
    for identity,items in groups.items():
        label=f'{identity[0]} 제{identity[1]}호증'+(f'의 {identity[2]}' if identity[2] else '')
        original=[f for f in old if f.confidence_features.get('exhibit')==label]
        if not original:continue
        tier='LEGACY_A_RETAINED';why='complete record/source/global/role uncertainty'
        cs=[certificates.get(i) for i,_,_ in items] if certificates else []
        if cs and all(cs) and all(t.literal(c['complete_title'])==t.literal(row['name']) for c,(_,row,_) in zip(cs,items)) and len({t.literal(c['complete_title']) for c in cs})>1:
            tier='CONFIRMED_TITLE_A';why='actual native full TITLE difference, dates preserved'
        elif all(r and r['source_complete'] and r['binding']['valid'] and all(r['native_date_occurrences']) and not row.get('in_section') for _,row,r in items):
            records=[r for _,_,r in items];inv=[r for r in records if r['heading']['inventory']];outside=[r for r in records if not r['heading']['inventory']]
            sameparent=len({r['heading']['parent'] for r in records})==1
            invclosed=all(r['heading']['end'] is not None and len(r['heading']['witnesses'])>=2 for r in inv)
            outsidebounded=all(r['heading']['end'] is not None or eof for r in outside)
            if sameparent and inv and outside and invclosed and outsidebounded and len({t.literal(r['complete_payload']) for r in inv})==1 and len({t.literal(r['complete_payload']) for r in outside})==1 and all(r['colon_shape'] for r in outside):
                tier='PENDING_CANDIDATE';why='whole record retained; date role/title role unproved; outside source extent bounded'
        tiers.append({'key':identity,'tier':tier,'reason':why,'original_A_count':len(original),
                      'records':[r and {'binding':r['binding'],'source_complete':r['source_complete'],'native_date_occurrences':r['native_date_occurrences'],'reference_span':r['reference_span'],'payload_span':r['payload_span'],'parent':r['heading']['parent'],'heading_closed':r['heading']['end'] is not None} for _,_,r in items]})
    assert s.exhibit_rows(doc)==before
    assert t.finding_contract(s.check_exhibits(doc))==t.finding_contract(findings)
    assert all(f.evidence_grade=='A' for f in old)
    return {'original_A_findings':len(old),'tiers':tiers,'source_EOF_proved':eof,'original_rows':len(rows),'original_rows_findings_unchanged':True,
            'other_original_rule_ids':[f.confidence_features.get('rule_id') for f in findings if f not in old]}


def run(directory):
    results=[]
    modes=('dated_closed','dated_eof','invalid_closed','invalid_eof','genuine_mixed_eof','protected_date','global_date','other_parent','missing_owner','quoted_eof','multi_date_eof','whole_date_only')
    for extension in ('txt','pdf','docx'):
        for number,(title,other) in enumerate((('점검기록','점검도면'),('작업대장','작업일지'),('관측도면','관측보고'))):
            for mode in modes:
                date='2025. 2. 30.' if mode.startswith('invalid') else '2025. 3. 4.'
                first=f'갑 제1호증: {title}, 확인 부분 {date} 경로 기록.'
                last=f'갑 제1호증 {other if mode=="genuine_mixed_eof" else title}: 순서를 설명한다. {date} 확인 기록.'
                lines=['1. 자료 제출','가. 공개 자료 목록',first,'을 제2호증: 비교자료, 확인 부분 2025. 3. 4. 경로 기록.','나. 자료 확인','기본 순서는 여러 단계로 구성되어 있다.','다. 처리 분석',last]
                if mode in ('dated_closed','invalid_closed'):lines+=['라. 별도 정리','첨부서류']
                if mode=='protected_date':lines[2]=f'갑 제1호증: {title} [2024. 1. 2.] 확인 부분';lines[7]=f'갑 제1호증 {other} [2024. 1. 3.]: 순서를 설명한다.'
                if mode=='global_date':lines.append(f'갑 제1호증 {title}: 순서를 설명한다. 2025. 3. 5. 확인 기록.')
                if mode=='other_parent':lines[6]='2. 별도 처리'
                if mode=='quoted_eof':lines=['《']+lines
                if mode=='multi_date_eof':lines[2]+=' 2026. 4. 5. 추가 확인.';lines[7]+=' 2026. 4. 5. 추가 확인.'
                if mode=='whole_date_only':lines[2]=f'갑 제1호증: {title} 2024. 1. 2.';lines[7]=f'갑 제1호증 {title} 2024. 1. 3.'
                folder=directory/f'{extension}_{number}_{mode}';folder.mkdir(parents=True,exist_ok=True)
                doc=s.parsed(folder,extension,lines);record=diagnose(doc,folder/('public_caption.'+extension),missing_owner=mode=='missing_owner')
                expected='PENDING_CANDIDATE' if mode in ('dated_closed','dated_eof','invalid_closed','invalid_eof','genuine_mixed_eof','multi_date_eof') else 'LEGACY_A_RETAINED'
                assert record['original_A_findings']==(0 if mode=='whole_date_only' else 1),(extension,mode,record)
                assert [i['tier'] for i in record['tiers']]==([] if mode=='whole_date_only' else [expected]),(extension,mode,record)
                if mode.startswith('invalid'):assert 'EVI.EVIDENCE_DATE_INVALID' in record['other_original_rule_ids']
                record.update(format=extension,mode=mode,example=number,public_true_duplicate=mode in ('genuine_mixed_eof','protected_date','global_date','whole_date_only'),missing_owner_simulation=mode=='missing_owner')
                results.append(record)
    for number,title in enumerate(('점검기록','작업대장','관측도면')):
        folder=directory/f'native_date_title_{number}';folder.mkdir(parents=True,exist_ok=True)
        names=[title+': 2024. 1. 2. 동쪽',title+': 2024. 1. 3. 서쪽']
        doc,_=t.v.table_document(folder,'pdf',[['갑 제1호증',name,'2025. 3. 4.','기재 확인'] for name in names])
        certs=t.native_title_certificates(doc,folder/'public_title_role.pdf')
        assert len(certs)==2
        record=diagnose(doc,folder/'public_title_role.pdf',certificates=certs)
        assert record['original_A_findings']==1 and [i['tier'] for i in record['tiers']]==['CONFIRMED_TITLE_A']
        record.update(format='pdf',mode='native_date_title',example=number,public_true_duplicate=True,actual_title_certificates=list(certs.values()))
        results.append(record)
    # Actual file counterexamples: complete rendered item mappings do not
    # certify EOF when other native source structure/order remains uncertain.
    for extension in ('pdf','docx'):
        for number,title in enumerate(('점검기록','작업대장','관측도면')):
            folder=directory/f'{extension}_source_boundary_{number}';folder.mkdir(parents=True,exist_ok=True)
            lines=['1. 자료 제출','가. 공개 자료 목록',f'갑 제1호증: {title}, 확인 부분 2025. 3. 4. 경로 기록.',
                   '을 제2호증: 비교자료, 확인 부분 2025. 3. 4. 경로 기록.','나. 자료 확인',
                   '기본 순서는 여러 단계로 구성되어 있다.','다. 처리 분석',
                   f'갑 제1호증 {title}: 순서를 설명한다. 2025. 3. 4. 확인 기록.']
            path=folder/('public_caption.'+extension)
            if extension=='pdf':
                from reportlab.pdfgen import canvas
                from reportlab.lib.pagesizes import A4
                s.build_pdf({'name':'public font','header':'공개','footer':'합성','body':[('p','공개')]},folder/'font.pdf')
                out=canvas.Canvas(str(path),pagesize=A4);out.setFont(s.build_pdf.__globals__['FONT'],10.5)
                # Visual reading order stays fixed; actual native drawing order
                # is reversed. No sorting it away for the EOF certificate.
                for index in reversed(range(len(lines))):out.drawString(50,750-20*index,lines[index])
                out.save();mode='native_source_reordered'
            else:
                from docx import Document
                out=Document()
                for line in lines:out.add_paragraph(line)
                out.add_table(rows=1,cols=1).cell(0,0).text='별도 표의 기재'
                out.save(path);mode='native_source_table'
            doc=s.parse_document(str(path),document_id='public_caption',filename=path.name,
                                 mime_type='application/pdf' if extension=='pdf' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',sha256='0'*64)
            record=diagnose(doc,path)
            assert record['original_A_findings']==1 and [i['tier'] for i in record['tiers']]==['LEGACY_A_RETAINED']
            assert not record['source_EOF_proved']
            record.update(format=extension,mode=mode,example=number,public_true_duplicate=False)
            results.append(record)
    counts=Counter()
    for record in results:
        counts['documents']+=1;counts['original_A']+=record['original_A_findings']
        if record['public_true_duplicate']:
            counts['public_true_documents']+=1
            if not record['original_A_findings']:counts['original_true_misses']+=1
        counts['source_EOF_proved']+=record['source_EOF_proved']
        counts['original_date_invalid_findings']+=record['other_original_rule_ids'].count('EVI.EVIDENCE_DATE_INVALID')
        for item in record['tiers']:
            counts[item['tier']]+=1
            if item['tier']=='PENDING_CANDIDATE':counts['true_pending' if record['public_true_duplicate'] else 'false_pending']+=1
    print(json.dumps({'records':results,'counts':dict(counts)},ensure_ascii=False,indent=2))


if __name__=='__main__':
    import sys
    run(Path(sys.argv[1]))
