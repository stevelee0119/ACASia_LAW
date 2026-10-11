"""Separate EOF safety comparison; frozen earlier diagnostics unchanged."""
import contextlib
import copy
import importlib.util
import io
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

spec=importlib.util.spec_from_file_location('dated',Path(__file__).with_name('TK-73_post013_dated_eof_spike.py'))
d=importlib.util.module_from_spec(spec);spec.loader.exec_module(d)
t=d.t;s=t.s


def source_support(path):
    """Small explicit supported-source domain; unknown objects are not empty."""
    reasons=[]
    if path.suffix=='.pdf':
        import pdfplumber
        with pdfplumber.open(path) as raw:
            for page in raw.pages:
                # Character-only native text scope. A graphic may contain text
                # although pdfplumber exposes no chars; no omission waiver.
                reasons.extend('unsupported_pdf_object:'+kind for kind,items in page.objects.items() if kind!='char' and items)
                if page.annots:reasons.append('unsupported_pdf_annotation')
    elif path.suffix=='.docx':
        from zipfile import ZipFile
        from xml.etree import ElementTree as ET
        with ZipFile(path) as archive:
            if any(re.match(r'word/(?:header|footer|comments|footnotes|endnotes)',name) for name in archive.namelist()):reasons.append('unsupported_word_source_part')
            tree=ET.fromstring(archive.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        body=tree.find('w:body',ns)
        if body is None:return ['missing_word_body']
        layout={'sectPr','pgSz','pgMar','cols','docGrid','type','titlePg','textDirection','vAlign','bidi','rtlGutter'}
        for child in body:
            tag=child.tag.rsplit('}',1)[-1]
            allowed={'p','r','t'} if tag=='p' else layout if tag=='sectPr' else set()
            for node in child.iter():
                kind=node.tag.rsplit('}',1)[-1]
                if not node.tag.startswith('{'+ns['w']+'}'):reasons.append('unsupported_word_namespace')
                if kind not in allowed:reasons.append('unsupported_word_node:'+kind)
    else:
        raw=path.read_text(encoding='utf-8')
        if any(unicodedata.category(c) in ('Cf','Co','Cs') or (unicodedata.category(c)=='Cc' and c not in '\t\r\n') for c in raw):reasons.append('unsupported_text_codepoint')
    return sorted(set(reasons))


def refined(doc,path,certificates=None,missing_owner=False):
    before_rows=copy.deepcopy(s.exhibit_rows(doc))
    before_findings=t.finding_contract(s.check_exhibits(doc))
    old=d.diagnose(doc,path,certificates=certificates,missing_owner=missing_owner)
    result=copy.deepcopy(old);text=s.build_reading_text(doc).text
    reasons=source_support(path)
    closed=not bool(s.protection(text)[-1])
    if any(unicodedata.category(c) in ('Cf','Co','Cs') or (unicodedata.category(c)=='Cc' and c not in '\t\r\n') for c in text):
        reasons.append('uncertain_reading_codepoint')
    support=not reasons
    # Do not conflate decoded-text sequence exhaustion with available-source
    # support or protection closure. No existing source failure is waived.
    result['c16_decoded_lexical_EOF']=old['source_EOF_proved']
    result['source_supported']=support;result['protection_closed']=closed
    result['source_EOF_proved']=old['source_EOF_proved'] and support and closed
    result['support_rejections']=reasons
    for item in result['tiers']:
        if item['tier']=='PENDING_CANDIDATE' and (not support or not closed):
            item['tier']='LEGACY_A_RETAINED'
            item['reason']='unsupported native source or trailing protection uncertainty; original A retained'
    assert s.exhibit_rows(doc)==before_rows
    assert t.finding_contract(s.check_exhibits(doc))==before_findings
    return old,result


def read_document(path):
    return s.parse_document(str(path),document_id='public_caption',filename=path.name,
                            mime_type={'.txt':'text/plain','.pdf':'application/pdf','.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}[path.suffix],sha256='0'*64)


def observation(old,new,mode,extension,number,truth):
    assert old['original_A_findings']==new['original_A_findings']
    assert old['original_rows']==new['original_rows']
    assert old['other_original_rule_ids']==new['other_original_rule_ids']
    return {'mode':mode,'format':extension,'example':number,'public_true_duplicate':truth,
            'original_A':old['original_A_findings'],'old_tiers':[i['tier'] for i in old['tiers']],
            'refined_tiers':[i['tier'] for i in new['tiers']],
            'decoded_lexical_EOF':new['c16_decoded_lexical_EOF'],'source_supported':new['source_supported'],
            'protection_closed':new['protection_closed'],'refined_source_EOF':new['source_EOF_proved'],
            'support_rejections':new['support_rejections'],'original_rows_findings_unchanged':True,
            'other_original_rule_ids':new['other_original_rule_ids']}


def new_controls(directory):
    records=[]
    for extension in ('txt','pdf','docx'):
        for number,(title,other) in enumerate((('점검기록','점검도면'),('작업대장','작업일지'),('관측도면','관측보고'))):
            for mode in ('ordinary_eof','genuine_mixed','global_date','balanced_trailer','trailing_open_quote'):
                lines=['1. 자료 제출','가. 공개 자료 목록',f'갑 제1호증: {title}, 확인 부분 2025. 3. 4. 경로 기록.',
                       '을 제2호증: 비교자료, 확인 부분 2025. 3. 4. 경로 기록.','나. 자료 확인',
                       '기본 순서는 여러 단계로 구성되어 있다.','다. 처리 분석',
                       f'갑 제1호증 {other if mode=="genuine_mixed" else title}: 순서를 설명한다. 2025. 3. 4. 확인 기록.']
                if mode=='global_date':lines.append(f'갑 제1호증 {title}: 순서를 설명한다. 2025. 3. 5. 확인 기록.')
                if mode=='balanced_trailer':lines.append('「추가 설명」')
                if mode=='trailing_open_quote':lines.append('「추가 설명')
                folder=directory/f'{extension}_{number}_{mode}';folder.mkdir(parents=True,exist_ok=True)
                doc=s.parsed(folder,extension,lines);path=folder/('public_caption.'+extension)
                old,new=refined(doc,path)
                assert old['original_A_findings']==1
                assert [i['tier'] for i in old['tiers']]==['LEGACY_A_RETAINED' if mode=='global_date' else 'PENDING_CANDIDATE']
                assert [i['tier'] for i in new['tiers']]==['LEGACY_A_RETAINED' if mode in ('global_date','trailing_open_quote') else 'PENDING_CANDIDATE']
                assert new['source_EOF_proved']==(mode!='trailing_open_quote')
                records.append(observation(old,new,mode,extension,number,mode in ('genuine_mixed','global_date')))
    for extension in ('pdf','docx'):
        for number,title in enumerate(('점검기록','작업대장','관측도면')):
            lines=['1. 자료 제출','가. 공개 자료 목록',f'갑 제1호증: {title}, 확인 부분 2025. 3. 4. 경로 기록.',
                   '을 제2호증: 비교자료, 확인 부분 2025. 3. 4. 경로 기록.','나. 자료 확인',
                   '기본 순서는 여러 단계로 구성되어 있다.','다. 처리 분석',
                   f'갑 제1호증 {title}: 순서를 설명한다. 2025. 3. 4. 확인 기록.']
            folder=directory/f'{extension}_native_unsupported_{number}';folder.mkdir(parents=True,exist_ok=True)
            path=folder/('public_caption.'+extension)
            if extension=='pdf':
                from PIL import Image
                from reportlab.pdfgen import canvas
                from reportlab.lib.pagesizes import A4
                s.build_pdf({'name':'public font','header':'공개','footer':'합성','body':[('p','공개')]},folder/'font.pdf')
                image=folder/'public_square.png';Image.new('RGB',(20,20),(15,30,45)).save(image)
                out=canvas.Canvas(str(path),pagesize=A4);out.setFont(s.build_pdf.__globals__['FONT'],10.5)
                for index,line in enumerate(lines):out.drawString(50,750-20*index,line)
                out.showPage();out.drawImage(str(image),50,650,width=60,height=60);out.save();mode='image_only_final_page'
            else:
                from docx import Document
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn
                out=Document()
                for line in lines:out.add_paragraph(line)
                symbol=OxmlElement('w:sym');symbol.set(qn('w:font'),'Wingdings');symbol.set(qn('w:char'),'F0FC')
                out.add_paragraph().add_run()._r.append(symbol);out.save(path);mode='native_symbol_at_end'
            old,new=refined(read_document(path),path)
            assert old['original_A_findings']==1 and [i['tier'] for i in old['tiers']]==['PENDING_CANDIDATE']
            assert old['source_EOF_proved'] and not new['source_EOF_proved'] and not new['source_supported']
            assert [i['tier'] for i in new['tiers']]==['LEGACY_A_RETAINED']
            records.append(observation(old,new,mode,extension,number,False))
    for number,title in enumerate(('점검기록','작업대장','관측도면')):
        from docx import Document
        from lxml import etree
        folder=directory/f'docx_foreign_namespace_{number}';folder.mkdir(parents=True,exist_ok=True)
        path=folder/'public_caption.docx';out=Document()
        for line in ['1. 자료 제출','가. 공개 자료 목록',f'갑 제1호증: {title}, 확인 부분 2025. 3. 4. 경로 기록.',
                     '을 제2호증: 비교자료, 확인 부분 2025. 3. 4. 경로 기록.','나. 자료 확인',
                     '기본 순서는 여러 단계로 구성되어 있다.','다. 처리 분석',
                     f'갑 제1호증 {title}: 순서를 설명한다. 2025. 3. 4. 확인 기록.']:out.add_paragraph(line)
        node=etree.Element('{urn:tk73-public-unknown}t');node.text='별도 원출처'
        out.add_paragraph().add_run()._r.append(node);out.save(path)
        old,new=refined(read_document(path),path)
        # The product reading includes this foreign node's text, while the
        # native w:t provider does not. Existing lexical binding already
        # abstains; namespace support is an independent rejection, not a gain.
        assert not old['source_EOF_proved'] and [i['tier'] for i in old['tiers']]==['LEGACY_A_RETAINED']
        assert not new['source_EOF_proved'] and 'unsupported_word_namespace' in new['support_rejections']
        assert [i['tier'] for i in new['tiers']]==['LEGACY_A_RETAINED']
        records.append(observation(old,new,'foreign_namespace_at_end','docx',number,False))
    return records


def frozen_comparisons(directory):
    with contextlib.redirect_stdout(io.StringIO()):original111=t.run(directory/'original111')
    capture=io.StringIO()
    with contextlib.redirect_stdout(capture):d.run(directory/'dated117')
    original117=json.loads(capture.getvalue())['records']
    groups={}
    for population,frozen in (('original111',original111),('dated117',original117)):
        records=[]
        for expected in frozen:
            mode=expected['mode'];extension=expected['format'];number=expected['example']
            native=mode in ('native_title','native_date_title')
            folder=(f'native_pdf_title_{number}' if mode=='native_title' else f'native_date_title_{number}' if mode=='native_date_title' else
                    f'{extension}_source_boundary_{number}' if mode in ('native_source_reordered','native_source_table') else f'{extension}_{number}_{mode}')
            path=directory/population/folder/('public_title_role.pdf' if native else 'public_caption.'+extension)
            doc=read_document(path);certs=t.native_title_certificates(doc,path) if native else None
            provider=s.source_index
            if population=='original111' and mode=='missing_owner':
                def original_fault(reading,pieces):
                    owners=provider(reading,pieces)
                    for index,owner in enumerate(owners):
                        if owner and reading.text[index]=='갑':owners[index]=None;break
                    return owners
                s.source_index=original_fault
            try:old,new=refined(doc,path,certificates=certs,missing_owner=population=='dated117' and mode=='missing_owner')
            finally:s.source_index=provider
            assert old['original_A_findings']==expected['original_A_findings']
            assert old['original_rows']==expected['original_rows']
            assert old['other_original_rule_ids']==expected['other_original_rule_ids']
            if population=='dated117':assert [i['tier'] for i in old['tiers']]==[i['tier'] for i in expected['tiers']]
            # Source-support fix must not change earlier generated tier results.
            assert [i['tier'] for i in new['tiers']]==[i['tier'] for i in old['tiers']]
            records.append(observation(old,new,mode,extension,number,expected['public_true_duplicate']))
        groups[population]=records
    return groups


def summary(records):
    counts=Counter()
    for record in records:
        counts['documents']+=1;counts['original_A']+=record['original_A'];counts['tier_changes']+=record['old_tiers']!=record['refined_tiers']
        counts['decoded_lexical_EOF']+=record['decoded_lexical_EOF'];counts['supported_source_EOF']+=record['refined_source_EOF']
        for tier in record['refined_tiers']:
            counts[tier]+=1
            if tier=='PENDING_CANDIDATE':counts['true_pending' if record['public_true_duplicate'] else 'false_pending']+=1
        counts['date_invalid_findings']+=record['other_original_rule_ids'].count('EVI.EVIDENCE_DATE_INVALID')
    return dict(counts)


def run(directory):
    new=new_controls(directory/'new_controls');frozen=frozen_comparisons(directory/'frozen')
    print(json.dumps({'new_controls':new,'frozen':frozen,'counts':{'new_controls':summary(new),**{name:summary(records) for name,records in frozen.items()}}},ensure_ascii=False,indent=2))


if __name__=='__main__':
    import sys
    run(Path(sys.argv[1]))
